from __future__ import annotations

import csv
import email
import json
import mimetypes
import re
import subprocess
import tempfile
import uuid
from dataclasses import dataclass, field
from email import policy
from pathlib import Path
from typing import Any

from bs4 import BeautifulSoup
from charset_normalizer import from_bytes
from docx import Document as DocxDocument
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation
from lxml import etree

from po_ocr.ingest import describe_file
import fitz
from po_ocr.pipeline import PurchaseOrderOCR
from po_ocr.extract import all_rows
from po_ocr.layout_guides import split_items_by_compartments

from .models import ContentBlock, EmbeddedObject, SourceRef, TableCell, UniversalTable


@dataclass
class ReadResult:
    raw_text: str = ""
    blocks: list[ContentBlock] = field(default_factory=list)
    tables: list[UniversalTable] = field(default_factory=list)
    embedded_objects: list[EmbeddedObject] = field(default_factory=list)
    pages: list[dict[str, Any]] = field(default_factory=list)
    structured_data: dict[str, Any] = field(default_factory=dict)
    page_results: list[Any] = field(default_factory=list)  # po_ocr PageResult objects
    family: str = "unknown"
    page_count: int | None = None
    sheet_count: int | None = None
    slide_count: int | None = None
    warnings: list[str] = field(default_factory=list)


def _block(kind: str, text: str, source: SourceRef, confidence: float = 1.0, **metadata: Any) -> ContentBlock:
    return ContentBlock(id=str(uuid.uuid4()), kind=kind, text=str(text), confidence=confidence, source=source, metadata=metadata)


def _table(name: str | None, data: list[list[Any]], source: SourceRef, confidence: float = 1.0) -> UniversalTable:
    rows = len(data)
    cols = max((len(r) for r in data), default=0)
    headers = [str(x) if x is not None else "" for x in (data[0] if data else [])]
    cells: list[TableCell] = []
    for r_idx, row in enumerate(data, start=1):
        for c_idx, value in enumerate(row, start=1):
            if value is None:
                continue
            cells.append(TableCell(row=r_idx, column=c_idx, text=str(value), confidence=confidence,
                                   source=SourceRef(page=source.page, sheet=source.sheet, slide=source.slide,
                                                    row=r_idx, column=c_idx)))
    return UniversalTable(id=str(uuid.uuid4()), name=name, source=source, rows=rows, columns=cols,
                          headers=headers, data=data, cells=cells, confidence=confidence)


def read_pdf_or_image(path: str | Path, po_engine: PurchaseOrderOCR) -> ReadResult:
    p = Path(path)
    fd = describe_file(p)
    pages, kinds = po_engine._local_pages(p, fd.extension)
    blocks: list[ContentBlock] = []
    page_payloads: list[dict[str, Any]] = []
    for page in pages:
        rows = all_rows([page], y_factor=float(po_engine.config["layout"].get("line_y_tolerance_factor", 0.65)))
        for row in rows:
            groups = split_items_by_compartments(
                row.words,
                page.vector_lines,
                page.vector_rectangles,
                gap_threshold=float(po_engine.config["layout"].get("visual_row_gap_threshold", 48.0)),
            )
            for group in groups:
                if not group["text"]:
                    continue
                blocks.append(_block(
                    "paragraph", group["text"], SourceRef(page=page.page, bbox=group["bbox"]), row.confidence,
                    source_type=page.source_type,
                    compartmentalized=len(groups) > 1,
                    boundary_source=group["boundary_source"],
                ))
        page_payloads.append({
            "page": page.page,
            "width": page.width,
            "height": page.height,
            "source_type": page.source_type,
            "confidence": page.confidence,
            "rotation_applied": page.rotation_applied,
            "text": page.text,
            "vector_layout": {
                "line_count": len(page.vector_lines),
                "rectangle_count": len(page.vector_rectangles),
            },
        })
    tables: list[UniversalTable] = []
    if fd.extension == "pdf":
        try:
            doc = fitz.open(p)
            for page_index, pdf_page in enumerate(doc, start=1):
                # PyMuPDF table extraction works best on native/vector PDFs.
                try:
                    finder = pdf_page.find_tables()
                    for t_idx, table in enumerate(finder.tables, start=1):
                        data = table.extract() or []
                        if data:
                            tables.append(_table(f"page_{page_index}_table_{t_idx}", data,
                                                 SourceRef(page=page_index), 0.98))
                except Exception:
                    pass
            doc.close()
        except Exception:
            pass
    return ReadResult(
        raw_text="\n\n".join(p.text for p in pages if p.text),
        blocks=blocks,
        tables=tables,
        pages=page_payloads,
        page_results=pages,
        family="pdf" if fd.extension == "pdf" else "image",
        page_count=len(pages),
    )


def read_docx(path: str | Path) -> ReadResult:
    doc = DocxDocument(path)
    blocks: list[ContentBlock] = []
    tables: list[UniversalTable] = []
    text_parts: list[str] = []
    for i, p in enumerate(doc.paragraphs, start=1):
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name or "").lower() if p.style else ""
        if "title" in style:
            kind = "title"
        elif "heading" in style:
            kind = "heading"
        elif "list" in style:
            kind = "list_item"
        else:
            kind = "paragraph"
        blocks.append(_block(kind, text, SourceRef(paragraph=i), 1.0, style=p.style.name if p.style else None))
        text_parts.append(text)

    for t_idx, tbl in enumerate(doc.tables, start=1):
        data = [[cell.text.strip() for cell in row.cells] for row in tbl.rows]
        tables.append(_table(f"table_{t_idx}", data, SourceRef(), 1.0))
        for row in data:
            text_parts.append(" | ".join(str(x) for x in row))

    # Headers and footers may contain critical order identifiers.
    for s_idx, section in enumerate(doc.sections, start=1):
        for kind, container in (("header", section.header), ("footer", section.footer)):
            for p in container.paragraphs:
                text = p.text.strip()
                if text:
                    blocks.append(_block(kind, text, SourceRef(), 1.0, section=s_idx))
                    text_parts.append(text)

    props = doc.core_properties
    structured = {
        "core_properties": {
            "title": props.title,
            "subject": props.subject,
            "author": props.author,
            "keywords": props.keywords,
            "comments": props.comments,
        }
    }
    return ReadResult(raw_text="\n".join(text_parts), blocks=blocks, tables=tables,
                      structured_data=structured, family="word")


def _cell_value(cell) -> Any:
    value = cell.value
    if value is None:
        return None
    if hasattr(value, "isoformat"):
        try:
            return value.isoformat()
        except Exception:
            pass
    return value


def read_xlsx(path: str | Path, max_cells_per_sheet: int = 100_000) -> ReadResult:
    wb = load_workbook(path, read_only=False, data_only=False)
    text_parts: list[str] = []
    blocks: list[ContentBlock] = []
    tables: list[UniversalTable] = []
    structured: dict[str, Any] = {"sheets": {}}
    warnings: list[str] = []

    for ws in wb.worksheets:
        max_row = ws.max_row or 0
        max_col = ws.max_column or 0
        if max_row * max_col > max_cells_per_sheet and max_row > 0:
            max_row = max(1, max_cells_per_sheet // max(1, max_col))
            warnings.append(f"Sheet {ws.title!r} truncated to {max_row} rows ({max_cells_per_sheet} cell cap).")
        data: list[list[Any]] = []
        formulas: dict[str, Any] = {}
        comments: dict[str, str] = {}
        for r in range(1, max_row + 1):
            row: list[Any] = []
            nonempty = False
            for c in range(1, max_col + 1):
                cell = ws.cell(r, c)
                value = _cell_value(cell)
                row.append(value)
                if value not in (None, ""):
                    nonempty = True
                    text_parts.append(str(value))
                    blocks.append(_block("cell", str(value), SourceRef(sheet=ws.title, row=r, column=c), 1.0,
                                         coordinate=cell.coordinate))
                    if isinstance(value, str) and value.startswith("="):
                        formulas[cell.coordinate] = value
                if cell.comment and cell.comment.text:
                    comments[cell.coordinate] = cell.comment.text
            if nonempty:
                data.append(row)
        if data:
            tables.append(_table(ws.title, data, SourceRef(sheet=ws.title), 1.0))
        structured["sheets"][ws.title] = {
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "merged_ranges": [str(x) for x in ws.merged_cells.ranges],
            "formulas": formulas,
            "comments": comments,
        }
    return ReadResult(raw_text="\n".join(text_parts), blocks=blocks, tables=tables,
                      structured_data=structured, family="spreadsheet", sheet_count=len(wb.worksheets), warnings=warnings)


def read_pptx(path: str | Path) -> ReadResult:
    prs = Presentation(path)
    text_parts: list[str] = []
    blocks: list[ContentBlock] = []
    tables: list[UniversalTable] = []
    objects: list[EmbeddedObject] = []
    pages: list[dict[str, Any]] = []

    for s_idx, slide in enumerate(prs.slides, start=1):
        slide_text: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = shape.text.strip()
                if text:
                    kind = "title" if shape == slide.shapes.title else "paragraph"
                    blocks.append(_block(kind, text, SourceRef(slide=s_idx), 1.0,
                                         left=int(shape.left), top=int(shape.top), width=int(shape.width), height=int(shape.height)))
                    slide_text.append(text)
                    text_parts.append(text)
            if getattr(shape, "has_table", False):
                data = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
                tables.append(_table(f"slide_{s_idx}_table", data, SourceRef(slide=s_idx), 1.0))
                for row in data:
                    text_parts.append(" | ".join(row))
            if shape.shape_type == 13:  # MSO_SHAPE_TYPE.PICTURE
                objects.append(EmbeddedObject(kind="image", name=getattr(shape, "name", None), source=SourceRef(slide=s_idx)))
        try:
            notes_text = "\n".join(p.text for p in slide.notes_slide.notes_text_frame.paragraphs if p.text.strip())
        except Exception:
            notes_text = ""
        if notes_text:
            blocks.append(_block("paragraph", notes_text, SourceRef(slide=s_idx), 1.0, notes=True))
            text_parts.append(notes_text)
        pages.append({"slide": s_idx, "text": "\n".join(slide_text)})
    return ReadResult(raw_text="\n".join(text_parts), blocks=blocks, tables=tables,
                      embedded_objects=objects, pages=pages, family="presentation", slide_count=len(prs.slides))


def _decode_text(path: str | Path) -> str:
    data = Path(path).read_bytes()
    best = from_bytes(data).best()
    if best is not None:
        return str(best)
    return data.decode("utf-8", errors="replace")


def read_delimited(path: str | Path, delimiter: str | None = None) -> ReadResult:
    text = _decode_text(path)
    sample = text[:8192]
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
        except Exception:
            delimiter = ","
    rows = list(csv.reader(text.splitlines(), delimiter=delimiter))
    blocks: list[ContentBlock] = []
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, value in enumerate(row, start=1):
            if value.strip():
                blocks.append(_block("cell", value, SourceRef(row=r_idx, column=c_idx), 1.0))
    table = _table(Path(path).stem, rows, SourceRef(), 1.0) if rows else None
    return ReadResult(raw_text=text, blocks=blocks, tables=[table] if table else [], family="data")


def read_text(path: str | Path) -> ReadResult:
    text = _decode_text(path)
    blocks: list[ContentBlock] = []
    for i, line in enumerate(text.splitlines(), start=1):
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("#"):
            kind = "heading"
        elif re.match(r"^(?:[-*+] |\d+[.)] )", stripped):
            kind = "list_item"
        else:
            kind = "paragraph"
        blocks.append(_block(kind, stripped, SourceRef(row=i), 1.0))
    return ReadResult(raw_text=text, blocks=blocks, family="text")


def read_json(path: str | Path) -> ReadResult:
    text = _decode_text(path)
    payload = json.loads(text)
    blocks = [_block("code", json.dumps(payload, ensure_ascii=False, indent=2), SourceRef(), 1.0, language="json")]
    return ReadResult(raw_text=text, blocks=blocks, structured_data={"json": payload}, family="data")


def read_xml(path: str | Path) -> ReadResult:
    data = Path(path).read_bytes()
    root = etree.fromstring(data)
    texts = [t.strip() for t in root.itertext() if t and t.strip()]
    blocks = [_block("paragraph", t, SourceRef(row=i), 1.0) for i, t in enumerate(texts, start=1)]
    return ReadResult(raw_text="\n".join(texts), blocks=blocks,
                      structured_data={"root_tag": root.tag}, family="markup")


def read_html(path: str | Path) -> ReadResult:
    text = _decode_text(path)
    soup = BeautifulSoup(text, "html.parser")
    blocks: list[ContentBlock] = []
    tables: list[UniversalTable] = []
    text_parts: list[str] = []
    for tag in soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "li"]):
        value = tag.get_text(" ", strip=True)
        if not value:
            continue
        kind = "heading" if tag.name.startswith("h") else "list_item" if tag.name == "li" else "paragraph"
        blocks.append(_block(kind, value, SourceRef(), 1.0, tag=tag.name))
        text_parts.append(value)
    for idx, tbl in enumerate(soup.find_all("table"), start=1):
        data = []
        for tr in tbl.find_all("tr"):
            data.append([c.get_text(" ", strip=True) for c in tr.find_all(["th", "td"])])
        if data:
            tables.append(_table(f"html_table_{idx}", data, SourceRef(), 1.0))
    return ReadResult(raw_text="\n".join(text_parts), blocks=blocks, tables=tables, family="markup")


def read_eml(path: str | Path) -> ReadResult:
    msg = email.message_from_bytes(Path(path).read_bytes(), policy=policy.default)
    blocks: list[ContentBlock] = []
    objects: list[EmbeddedObject] = []
    parts: list[str] = []
    for key in ("Subject", "From", "To", "Cc", "Date", "Message-ID"):
        value = msg.get(key)
        if value:
            line = f"{key}: {value}"
            blocks.append(_block("key_value", line, SourceRef(), 1.0, key=key, value=str(value)))
            parts.append(line)
    for part in msg.walk():
        ctype = part.get_content_type()
        filename = part.get_filename()
        if filename:
            data = part.get_payload(decode=True) or b""
            objects.append(EmbeddedObject(kind="attachment", name=filename, mime_type=ctype, size_bytes=len(data)))
            continue
        if ctype == "text/plain":
            try:
                body = part.get_content()
            except Exception:
                body = ""
            if body:
                blocks.append(_block("paragraph", body.strip(), SourceRef(), 1.0))
                parts.append(body)
        elif ctype == "text/html" and not any(p.get_content_type() == "text/plain" for p in msg.walk()):
            try:
                html = part.get_content()
                body = BeautifulSoup(html, "html.parser").get_text("\n", strip=True)
            except Exception:
                body = ""
            if body:
                blocks.append(_block("paragraph", body, SourceRef(), 1.0))
                parts.append(body)
    return ReadResult(raw_text="\n".join(parts), blocks=blocks, embedded_objects=objects, family="data")


def convert_with_libreoffice(path: str | Path, target_ext: str = "pdf") -> Path:
    src = Path(path)
    temp_dir = Path(tempfile.mkdtemp(prefix="uda_lo_"))
    proc = subprocess.run([
        "libreoffice", "--headless", "--convert-to", target_ext,
        "--outdir", str(temp_dir), str(src)
    ], stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, timeout=120)
    if proc.returncode != 0:
        raise ValueError(f"LibreOffice conversion failed: {proc.stderr.strip() or proc.stdout.strip()}")
    outputs = list(temp_dir.glob(f"*.{target_ext}"))
    if not outputs:
        raise ValueError(f"LibreOffice did not produce a .{target_ext} output")
    return outputs[0]


def read_any(path: str | Path, po_engine: PurchaseOrderOCR, max_cells_per_sheet: int = 100_000) -> ReadResult:
    p = Path(path)
    ext = p.suffix.lower().lstrip(".")
    if ext in {"pdf", "png", "jpg", "jpeg", "tif", "tiff", "bmp", "webp"}:
        return read_pdf_or_image(p, po_engine)
    if ext == "docx":
        return read_docx(p)
    if ext in {"xlsx", "xlsm"}:
        return read_xlsx(p, max_cells_per_sheet=max_cells_per_sheet)
    if ext == "pptx":
        return read_pptx(p)
    if ext == "csv":
        return read_delimited(p, delimiter=None)
    if ext == "tsv":
        return read_delimited(p, delimiter="\t")
    if ext in {"txt", "md", "log", "ini", "cfg", "conf", "yaml", "yml"}:
        return read_text(p)
    if ext == "json":
        return read_json(p)
    if ext == "xml":
        return read_xml(p)
    if ext in {"html", "htm"}:
        return read_html(p)
    if ext == "eml":
        return read_eml(p)
    if ext in {"doc", "xls", "ppt", "odt", "ods", "odp", "rtf"}:
        pdf = convert_with_libreoffice(p, "pdf")
        rr = read_pdf_or_image(pdf, po_engine)
        rr.family = "legacy_office"
        rr.warnings.append(f"{ext.upper()} read through LibreOffice PDF conversion.")
        return rr
    raise ValueError(f"Unsupported file extension: .{ext}")
